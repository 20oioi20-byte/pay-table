"""M365 연결 테스트용 가상 급여테이블 엑셀 만들기 (금액·센터명 모두 가상).

데모 HTML의 샘플(tools/demo_inject.js samplePayload)과 같은 행·열 배치로
'급여테이블(변경후)'·'급여테이블(변경전)' 두 시트를 만든다. 합계 행은 수식으로 넣어
v2 스크립트의 수식(formulas) 읽기도 함께 시험할 수 있다.

    python3 tools/make_test_xlsx.py
"""
from pathlib import Path
import openpyxl
from openpyxl.utils import get_column_letter as L

OUT = Path(__file__).resolve().parent.parent / 'samples' / '테스트용_가상급여테이블.xlsx'
ITEMS = ['기본급', '직책수당', '직무수당', '근속수당', '급식비', '통근비']
GRADE_ADD = [150000, 100000, 60000, 30000, 0, 0]
MAIN = [('일반상담', '사원', 2150000, 0, 80000, 30000, 150000, 0),
        ('일반상담', '주임', 2250000, 0, 80000, 50000, 150000, 0),
        ('일반상담', '선임', 2350000, 0, 80000, 70000, 150000, 0),
        ('신입', '수습', 2000000, 0, 60000, 0, 0, 0),
        ('야간상담', '심야상담', 2200000, 0, 100000, 30000, 0, 220000)]
LEAD = [('리더', '팀장', 2700000, 300000, 120000, 80000, 300000, 0),
        ('리더', '센터장', 3300000, 600000, 150000, 100000, 0, 0)]


def sheet(ws, raise_rate):
    ws['A1'] = '2026년 급여테이블 — 테스트용 가상 데이터'
    ws['A3'] = '센터명'; ws['D3'] = '테스트 가상센터'

    def block(title, row, jobs):
        ws.cell(row, 1, title)
        last = None
        for i, (g, n, base, pos, duty, ten, ot, night) in enumerate(jobs):
            c = 4 + i; col = L(c)
            if g != last:
                ws.cell(row + 1, c, g); last = g
            ws.cell(row + 2, c, n)
            vals = [round(base * raise_rate, -1), pos, duty, ten, 200000, 0]
            for k, lab in enumerate(ITEMS):
                ws.cell(row + 3 + k, 1, lab); ws.cell(row + 3 + k, c, vals[k])
            ws.cell(row + 9, 1, '기준급여 합계')
            ws.cell(row + 9, c, f'=SUM({col}{row + 3}:{col}{row + 8})')
            for k, g2 in enumerate(['A.', 'B.', 'C.', 'D.', 'E.', 'F.']):
                ws.cell(row + 10 + k, 3, g2)
                ws.cell(row + 10 + k, c, f'={col}{row + 9}+{GRADE_ADD[k]}')
            ws.cell(row + 16, 1, '초과근무수당'); ws.cell(row + 16, c, ot)
            ws.cell(row + 17, 1, '야간근로수당'); ws.cell(row + 17, c, night)
            ws.cell(row + 18, 1, '통신비'); ws.cell(row + 18, c, 30000)
            ws.cell(row + 19, 1, '월 총액')
            ws.cell(row + 19, c, f'={col}{row + 9}+SUM({col}{row + 16}:{col}{row + 18})')
        ws.cell(row + 10, 19, '업적평가 등급별 월 총액(기준급여 합계 + 업적수당)')
        ws.cell(row + 16, 19, '사원 10H, 팀장 20H')
        ws.cell(row + 19, 19, '통신비 포함')

    block('Ⅰ. 보수테이블', 5, MAIN)
    block('Ⅱ. 리딩그룹 보수테이블', 27, LEAD)
    ws['A50'] = '예외 및 특례사항'
    ws['A51'] = '1. 근무형태'
    for i, h in enumerate(['직무', '호봉', '근무일', '', '근무시간', '휴게시간', '비고']):
        ws.cell(52, 1 + i, h)
    rows = [['일반상담', '1~5호봉', '월~금', '', '09:00~18:00', '12:00~13:00', '주 40시간'],
            ['심야상담', '1~3호봉', '일~목', '', '22:00~07:00', '02:00~03:00', '야간수당 별도'],
            ['리더', '-', '월~금', '', '08:30~17:30', '12:00~13:00', '']]
    for k, r in enumerate(rows):
        for i, x in enumerate(r):
            ws.cell(53 + k, 1 + i, x)
    ws['A58'] = '2. 수습기간(3개월) 중에는 수습 기준표를 적용하고, 종료 후 사원 1호봉으로 전환한다.'
    ws['A60'] = '3. 초과근무수당은 직무별 약정시간 기준이며, 실제 근무시간이 초과할 경우 별도 정산한다.'
    ws['A63'] = '10. 기타'


def main():
    wb = openpyxl.Workbook()
    before = wb.active; before.title = '급여테이블(변경전)'
    sheet(before, 0.97)
    sheet(wb.create_sheet('급여테이블(변경후)'), 1.0)
    OUT.parent.mkdir(exist_ok=True)
    wb.save(OUT)
    print(OUT.name)


if __name__ == '__main__':
    main()
