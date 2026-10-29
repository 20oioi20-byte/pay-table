"""실제 급여테이블과 행·열 구조가 같은 가상 엑셀 만들기 (센터명·담당자·금액 모두 가상).

실제 양식(보수테이블 32~61행, 리딩그룹 63~92행, 예외 및 특례사항 94행~)과 같은 위치에
같은 항목명·수식 형태를 넣는다. 파서·검증 엔진 회귀 테스트와 M365 연결 시험에 쓴다.

    python3 tools/make_real_layout_xlsx.py
"""
from pathlib import Path
import shutil
import subprocess
import tempfile
import openpyxl
from openpyxl.utils import get_column_letter as L

OUT = Path(__file__).resolve().parent.parent / 'samples' / '가상_실양식_급여테이블.xlsx'
MW = 10320                      # 2026년 최저시급
MAIN = [('수습', 0, 150000), ('상담1', 0, 150000), ('상담2', 20000, 150000), ('상담3', 40000, 150000)]
LEAD = [('D', '팀장', 200000, 150000, 50000), ('E', '강사', 150000, 100000, 30000), ('Q', '센터장', 400000, 300000, 80000)]
GRADES = ['S', 'A', 'B', 'C.', 'D', 'E']


def header(ws, center, pm):
    ws['A1'] = f'{center} 직무GROUP 인력운영 기준'
    ws['A4'] = 'Ⅰ. 관리'; ws['S4'] = f'[ 담당PM : {pm} ]'
    ws['A6'] = '요청조직명'; ws['D6'] = '가상사업본부 가상사업단 가상1팀'; ws['H6'] = '요청일자'; ws['K6'] = 46023
    ws['A7'] = '요청사유'; ws['D7'] = '2026년 최저임금 반영'; ws['H7'] = '관리번호'; ws['K7'] = f'{center}/2026년1월'
    ws['A9'] = 'Ⅱ. 직무Group'
    ws['A11'] = '직군명'; ws['D11'] = 'TM직군'
    ws['A12'] = '직무Group명'; ws['D12'] = center; ws['H12'] = '직무Group정의'; ws['K12'] = '고객상담'
    ws['A14'] = 'Ⅲ. 직무'; ws['R14'] = '(단위 : 원, 명)'
    for c, v in zip('ABEHIRS', ['구분', '직무명', '직무정의', '직무/직책수당', '직무규정(소속)', '비평가시 / (등급)', '정원']):
        ws[f'{c}15'] = v
    rows = [('조직장', '센터장', '=Q67'), (None, '팀장', '=D67'), (None, '교육/QAA', '=E67'), ('조직원', '상담사', None)]
    for i, (a, b, h) in enumerate(rows):
        r = 16 + i
        if a: ws[f'A{r}'] = a
        ws[f'B{r}'] = b; ws[f'E{r}'] = f'{b} 업무'; ws[f'R{r}'] = 'X'; ws[f'S{r}'] = 1 if i < 3 else 20
        if h: ws[f'H{r}'] = h
    ws['A22'] = '계'; ws['S22'] = '=SUM(S16:S21)'
    ws['A24'] = 'Ⅳ. 정원'; ws['R24'] = '(단위 : 명)'
    ws['A25'] = '구분'; ws['B25'] = '직급'; ws['D25'] = '컨설턴트(일반)'; ws['I25'] = '컨설턴트(전문)'; ws['S25'] = '계'
    ws['B26'] = '호봉'
    for i, (n, *_r) in enumerate(MAIN):
        ws.cell(26, 4 + i, n); ws.cell(27, 4 + i, [3, 8, 6, 3][i])
    ws['A27'] = '정원'; ws['S27'] = '=SUM(D27:R27)'
    ws['A28'] = '구분'; ws['B28'] = '직급'; ws['D28'] = 'Staff'; ws['P28'] = 'Leader'; ws['S28'] = '계'
    ws['B29'] = '호봉'; ws['D29'] = '상담팀장'; ws['E29'] = '교육강사'; ws['R29'] = '센터장'
    ws['A30'] = '정원'; ws['D30'] = 2; ws['E30'] = 1; ws['R30'] = 1; ws['S30'] = '=SUM(D30:R30)'
    ws['P31'] = '총계'; ws['S31'] = '=S27+S30'


def pay_block(ws, top, title, cols, rate, lead):
    """top = 제목 행(32 또는 63). 항목은 top+3 ~ top+8, 소계 top+9 …"""
    ws.cell(top, 1, title); ws.cell(top, 18, '(단위 : 원, 명)')
    ws.cell(top + 1, 1, '구분'); ws.cell(top + 1, 19, '비고')
    if lead:
        ws.cell(top + 1, 4, 'Staff'); ws.cell(top + 1, 16, 'Leader')
    else:
        ws.cell(top + 1, 4, '상담')
    items = ['기본급', '직책수당', '직무수당', '근속수당', '급식비', '통근비']
    for k, lab in enumerate(items):
        ws.cell(top + 3 + k, 1, lab)
    r_sub, r_g, r_ot, r_ng, r_tel, r_tot = top + 9, top + 10, top + 16, top + 17, top + 18, top + 19
    ws.cell(r_sub, 1, '통상임금소계'); ws.cell(r_g, 1, '업적수당')
    for k, g in enumerate(GRADES):
        ws.cell(r_g + k, 3, g)
    ws.cell(r_ot, 1, '초과근무수당'); ws.cell(r_ng, 1, '야간근로수당'); ws.cell(r_tel, 1, '통신비')
    ws.cell(r_tot, 1, '월 보수 합계')
    labels = ['연차수당(1D)', '시간외수당(1H)', '휴일근로수당(1H)', '야근수당(1H)', '호칭별 연간 보수 합계',
              '전체 연간 인건비 합계(정원기준)', '당해년도 입사자 국민연금 산정 기준 보수월액',
              '당해년도 입사자 건강/고용/산재보험 산정 기준 보수월액']
    for k, lab in enumerate(labels):
        ws.cell(r_tot + 1 + k, 1, lab)
    ws.cell(r_tot + 10, 1, '변경 전 월보수 GAP')
    ws.cell(r_g, 19, '평가결과에 따름 '); ws.cell(r_tot, 19, 'B등급 기준'); ws.cell(r_tot + 5, 19, 'B등급 기준')
    ws.cell(r_tot + 7, 19, 'D등급 기준'); ws.cell(r_tot + 8, 19, 'B등급 기준')
    ws.cell(r_tot + 6, 19, f'=SUM(D{r_tot + 6}:R{r_tot + 6})')
    for c in range(4, 19):
        X = L(c)
        ws.cell(r_sub, c, f'=SUM({X}{top + 3}:{X}{top + 8})')
        ws.cell(r_tot, c, f'={X}{r_sub}+{X}{r_g + 2}+SUM({X}{r_ot}:{X}{r_tel})')
        for k, (lab, mul) in enumerate([('연차', '8'), ('시간외', '1.5'), ('휴일', '1.5'), ('야근', '0.5')]):
            ws.cell(r_tot + 1 + k, c, f'={X}{r_sub}/209*{mul}')
        ws.cell(r_tot + 5, c, f'={X}{r_tot}*12')
        ws.cell(r_tot + 6, c, f'={X}{r_tot + 5}*{X}{27 if not lead else 30}')
        ws.cell(r_tot + 7, c, f'={X}{r_sub}+{X}{r_g + 4}+{X}{r_ot}+{X}{r_ng}+{X}{r_tel}+$C$62-{X}{top + 7}')
        ws.cell(r_tot + 8, c, f'={X}{r_sub}+{X}{r_g + 2}+{X}{r_ot}+{X}{r_ng}+{X}{r_tel}+$C$62')
        ws.cell(r_g + 5, c, 0)                                    # E등급 0원
    if lead:
        ws.cell(r_ot, 19, '팀장, 강사, 센터장 5H')
        for col, name, pos, duty, ten in cols:
            c = openpyxl.utils.column_index_from_string(col); X = col
            ws.cell(top + 2, c, name)
            ws.cell(top + 3, c, f'={MW}*209')
            for k, v in enumerate([pos, duty, ten, round(100000 * rate, -3)]):
                ws.cell(top + 4 + k, c, v)
            for k, add in enumerate([200000, 150000, 100000, 50000, 20000]):
                ws.cell(r_g + k, c, add)
            ws.cell(r_ot, c, f'=ROUNDUP({X}{r_tot + 2}*5,-1)')
            ws.cell(r_tot + 10, c, f"={X}{r_tot}-'급여테이블(변경전)'!{X}{r_tot}")
    else:
        ws.cell(r_ot, 19, '2H')
        for i, (name, ten, _) in enumerate(cols):
            c = 4 + i; X = L(c)
            ws.cell(top + 2, c, name)
            ws.cell(top + 3, c, f'={MW}*209')
            if ten: ws.cell(top + 6, c, ten)
            ws.cell(top + 7, c, round(100000 * rate, -3))
            for k, add in enumerate([120000, 90000, 60000, 30000, 10000]):
                ws.cell(r_g + k, c, add)
            if i: ws.cell(r_ot, c, f'=ROUNDUP({X}{r_tot + 2}*2,-1)')
            ws.cell(r_tot + 10, c, f"={X}{r_tot}-'급여테이블(변경전)'!{X}{r_tot}")


def exceptions(ws):
    ws['A94'] = 'Ⅵ. 예외 및 특례사항'
    lines = {95: '     1. 수습기간은 3개월 적용한다.', 98: '     2. 무급사유 발생 시 통상임금 항목은 일할 계산한다.',
             99: '     3. 업적수당 : 성과평가 결과에 따라 직무별 차등 지급한다. ', 100: '     4. 복리후생비 지급 기준',
             101: '     ㅇ 급식비 ', 102: '구분', 103: '지급기준', 104: '제외사항 ', 113: '  5. 인센티브 지급 기준 ',
             117: '  6. 호봉별 등급부여 기준', 122: '   7. 연차 평균 사용률', 125: '   8. 월 근무일수에 따른 업적수당 지급 기준',
             128: '  9.근무형태'}
    for r, t in lines.items():
        ws[f'A{r}'] = t
    ws['B102'] = '주요내용 '; ws['B103'] = '매 월 마지막 영업일 재직 시 전액 지급'; ws['B104'] = '실제 출근일 수 5일 이하자는 지급 제외'
    for c, v in zip('ABCEFG', ['직무', ' 호봉', ' 근무일', ' 근무시간', ' 휴게시간', ' 비고']):
        ws[f'{c}129'] = v
    for k, (a, b) in enumerate([('센터장', '센터장'), ('팀장, 강사', '팀장, 강사'), ('상담사', '상담사')]):
        r = 130 + k
        ws[f'A{r}'] = a; ws[f'B{r}'] = b; ws[f'C{r}'] = '주 5일(월~금), 40시간(일 8H)'
        ws[f'E{r}'] = '09:00 ~ 18:00'; ws[f'F{r}'] = ' 12:00 ~ 13:00'


def sheet(ws, rate):
    header(ws, '가상시 민원상담콜센터', '홍길동')
    pay_block(ws, 32, 'Ⅴ. 보수테이블', MAIN, rate, False)
    ws['C62'] = 0
    ws['A63'] = 'ㅇ 리딩그룹(스탭이상 ~ 센터장) '
    pay_block(ws, 63, 'ㅇ 리딩그룹(스탭이상 ~ 센터장) ', LEAD, rate, True)
    exceptions(ws)


# ---------- v2.6: 견적서 · 견적대비 비교 (실제 파일과 같은 행·열 위치, 금액은 가상) ----------
CMP_JOBS = [('센터장', 'Q', 72, 1, [0, 1, 0, 0, 0, 0]), ('팀장', 'D', 72, 2, [0, 1, 1, 0, 0, 0]), ('강사', 'E', 72, 1, [0, 0, 1, 0, 0, 0]),
            ('수습', 'D', 41, 3, [0, 0, 0, 0, 0, 3]), ('상담1', 'E', 41, 8, [1, 2, 3, 1, 1, 0]), ('상담2', 'F', 41, 6, [1, 1, 3, 1, 0, 0]),
            ('상담3', 'G', 41, 3, [0, 1, 2, 0, 0, 0])]


def estimate(ws, center):
    ws['B2'] = '○ 센터명 :'; ws['C2'] = center; ws['B3'] = '○ 계약기간 :'; ws['C3'] = '2026.01.01. ~ 2026.12.31.'
    ws['J4'] = '(단위 : 원)'
    for c, v in zip('BEFGHIJ', ['구분', '센터장', '팀장/강사', '일반상담', '심야상담', '계', '비고(산출근거)']):
        ws[f'{c}5'] = v
    ws['B6'] = '인원 ①'; ws['E6'] = 1; ws['F6'] = 3; ws['G6'] = 20; ws['I6'] = '=SUM(E6:H6)'; ws['J6'] = '계약 정원 기준'
    ws['B7'] = '인건비'; ws['C7'] = '직접 인건비'
    rows = [('기본급', [2160000, 2160000, 2160000]), ('초과수당', [130000, 100000, 40000]), ('연차수당', None),
            ('식대/교통비', [150000, 150000, 150000]), ('직책/전문수당', [700000, 300000, None]),
            ('성과급(인센티브)', [150000, 100000, 60000]), ('휴일근로수당', [0, 0, 0])]
    for k, (lab, vals) in enumerate(rows):
        r = 7 + k; ws[f'D{r}'] = lab; ws[f'I{r}'] = f'=SUMPRODUCT(E{r}:H{r},$E$6:$H$6)'
        for c, v in zip('EFG', vals or []):
            if v is not None: ws[f'{c}{r}'] = v
    ws['D14'] = '직접인건비 소계'
    for c in 'EFGH': ws[f'{c}14'] = f'=SUM({c}7:{c}13)'
    ws['I14'] = '=SUMPRODUCT(E14:H14,$E$6:$H$6)'; ws['J14'] = '실비 청구 항목 등 표기'
    ws['C15'] = '간접 인건비'; ws['D15'] = '퇴직충당금'; ws['J15'] = '(직접인건비)/12'
    for c in 'EFG': ws[f'{c}15'] = f'=ROUND({c}14/12,0)'
    ws['D26'] = '간접인건비 소계'; ws['B36'] = '일반관리비 ③ 값 반영⇒'; ws['B37'] = '기업이윤 ④   값 반영⇒'; ws['B40'] = '총합계(vat 포함)'
    for c in 'EFG': ws[f'{c}34'] = 30000


def compare(ws):
    S = "'급여테이블(변경후)'!"
    ws['B1'] = '급여테이블 견적 대사'; ws['B3'] = '* ERP상 직접비 총계(직전년도연간 인건비)'
    ws['B5'] = '구분(연간)'; ws['J5'] = '급여 테이블'; ws['J6'] = '호봉'; ws['B7'] = '직접비'; ws['C7'] = '직접비 총계'; ws['F7'] = 0.03
    ws['J7'] = '호봉별 총원'; ws['J8'] = '통상임금'; ws['J11'] = '초과근로/야근수당'; ws['J12'] = '소계'
    ws['J13'] = '업적수당'; ws['K13'] = '등급'; ws['K21'] = '계'; ws['J22'] = '월 보수 합계'
    for k, g in enumerate(['S', 'A', 'B', 'C', 'D', 'E', 'F']):
        ws.cell(14 + k, 11, g)
    for i, (name, col, sub, n, dist) in enumerate(CMP_JOBS):
        c = 12 + i * 2; X, Y = L(c), L(c + 1)
        ws.cell(6, c, name); ws.cell(7, c, n); ws.cell(13, c, '단가'); ws.cell(13, c + 1, '인원')
        ws.cell(8, c, f'={S}{col}{sub}')
        ws.cell(11, c, f'={S}{col}{sub + 7}')
        ws.cell(12, c, f'=SUM({X}8:{Y}11)*{X}7')
        for k in range(6):
            ws.cell(14 + k, c, f'={S}{col}{sub + 1 + k}'); ws.cell(14 + k, c + 1, dist[k])
        ws.cell(21, c, f'=SUMPRODUCT({X}14:{X}20,{Y}14:{Y}20)')
        ws.cell(22, c, f'={X}12+{X}21')
    last = L(12 + len(CMP_JOBS) * 2 - 1)
    ws['K7'] = f'=SUM(L7:{last}7)'; ws['K12'] = f'=SUM(L12:{last}12)'; ws['K22'] = f'=SUM(L22:{last}22)'
    ws['B17'] = '<표2>'; ws['B18'] = '구분'; ws['D18'] = '견적'; ws['B19'] = '직접비'
    for c, v in zip('DEFGH', ['센터장', '팀장/강사', '상담사', '심야상담', '총계']):
        ws[f'{c}19'] = v
    ws['C20'] = '인원'; ws['D20'] = 1; ws['E20'] = 3; ws['F20'] = 20; ws['H20'] = '=SUM(D20:G20)'
    ws['C21'] = '직접비 단가'
    for c, e in zip('DEFG', 'EFGH'): ws[f'{c}21'] = f'=견적서!{e}14'
    ws['C22'] = '직무별 직접비 계'
    for c in 'DEFG': ws[f'{c}22'] = f'={c}20*{c}21'
    ws['H22'] = '=SUM(D22:G22)'
    ws['C23'] = '월 보수 가용 범위'; ws['G23'] = '=1-F7'; ws['H23'] = '=H22*G23'
    ws['C24'] = 'Gap'; ws['H24'] = '=H23-K22'


def main():
    wb = openpyxl.Workbook()
    before = wb.active; before.title = '급여테이블(변경전)'
    sheet(before, 0.9)
    sheet(wb.create_sheet('급여테이블(변경후)'), 1.0)
    estimate(wb.create_sheet('견적서'), '가상시 민원상담콜센터')
    compare(wb.create_sheet('견적대비 비교'))
    OUT.parent.mkdir(exist_ok=True)
    wb.save(OUT)
    recalc(OUT)
    print(OUT.name)


def recalc(path):
    """openpyxl은 수식 결과값을 넣지 못하므로 LibreOffice로 한 번 계산해 값까지 저장한다(없으면 건너뜀)."""
    soffice = shutil.which('soffice')
    if not soffice:
        print('LibreOffice가 없어 수식 결과값 없이 저장했습니다. Excel에서 한 번 열어 저장하면 값이 채워집니다.')
        return
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run([soffice, '--headless', '--calc', '--convert-to', 'xlsx:Calc MS Excel 2007 XML', '--outdir', tmp, str(path)],
                       check=True, capture_output=True, timeout=180)
        shutil.move(str(Path(tmp) / path.name), path)


if __name__ == '__main__':
    main()
