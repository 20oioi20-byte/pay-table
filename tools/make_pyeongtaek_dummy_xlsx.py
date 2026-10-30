"""실제와 비슷한 더미 급여테이블(예시: 평택) 만들기 — 변경동의서·조견표 추출 확인용.

실제 양식과 같은 행 위치(보수테이블 32행~, 리딩그룹 63행~, 예외 및 특례사항 94행~)에
실제 센터에서 흔한 구성을 넣는다. 금액·이름은 모두 더미.
  - 상담 호봉 12개(일반상담 5 · 전문상담 3 · 야간상담 2 · 단축상담 1 · 채팅상담 1) + 리딩그룹 5개
  - 변경전 = 2025년 최저시급 10,030원, 변경후 = 2026년 10,320원 (기본급 = 시급 × 소정시간 수식)
  - 급식비 150,000 → 200,000, 전문상담 업적수당 S 180,000 → 200,000, 팀장 직책수당 200,000 → 250,000
  - 단축상담: 주 25시간(일 5H) → 주 20시간(일 4H) 근무시간 변경
  - 채팅상담 · 사원: 변경후에만 있는 신규 호봉 (변경동의서 제외 대상)
  - 관리번호 '평택고객센터/2026년1월' → 시행일 2026. 01. 01.

    python3 tools/make_pyeongtaek_dummy_xlsx.py
"""
from pathlib import Path
import openpyxl
from openpyxl.utils import get_column_letter as L
from make_real_layout_xlsx import recalc

OUT = Path(__file__).resolve().parent.parent / 'samples' / '더미_평택_급여테이블.xlsx'
CENTER = '평택고객센터'
GRADES = ['S', 'A', 'B', 'C.', 'D', 'E']

# (그룹, 호봉, 월 소정시간, 직무수당, 근속수당, 업적 S~D, 초과 H, 야간 H)
NORMAL = [150000, 110000, 80000, 40000, 10000]
EXPERT = {'before': [180000, 130000, 90000, 50000, 20000], 'after': [200000, 140000, 100000, 50000, 20000]}
NIGHT = [160000, 120000, 90000, 40000, 10000]
SHORT = [70000, 50000, 30000, 10000, 0]


def main_jobs(when):
    jobs = [('일반상담', '수습', 209, 0, 0, NORMAL, 0, 0), ('일반상담', '사원', 209, 0, 0, NORMAL, 4, 0),
            ('일반상담', '주임', 209, 0, 20000, NORMAL, 4, 0), ('일반상담', '선임', 209, 0, 40000, NORMAL, 4, 0),
            ('일반상담', '책임', 209, 0, 60000, NORMAL, 4, 0),
            ('전문상담', '사원', 209, 30000, 0, EXPERT[when], 4, 0), ('전문상담', '주임', 209, 30000, 20000, EXPERT[when], 4, 0),
            ('전문상담', '선임', 209, 30000, 40000, EXPERT[when], 4, 0),
            ('야간상담', '사원', 209, 50000, 0, NIGHT, 0, 40), ('야간상담', '선임', 209, 50000, 40000, NIGHT, 0, 40),
            ('단축상담', '사원', 130 if when == 'before' else 104, 0, 0, SHORT, 0, 0)]
    if when == 'after':
        jobs.append(('채팅상담', '사원', 209, 20000, 0, NORMAL, 4, 0))
    return jobs


def lead_jobs(when):
    team = 200000 if when == 'before' else 250000
    # (그룹, 호봉, 직책수당, 직무수당, 근속수당, 초과 H)
    return [('Staff', '파트장', 100000, 100000, 30000, 5), ('Staff', '팀장', team, 150000, 50000, 5),
            ('Staff', '강사', 100000, 100000, 30000, 5), ('Leader', '실장', 300000, 200000, 60000, 5),
            ('Leader', '센터장', 500000, 300000, 80000, 0)]


def header(ws, jobs, leads):
    ws['A1'] = f'{CENTER} 직무GROUP 인력운영 기준'
    ws['A4'] = 'Ⅰ. 관리'; ws['S4'] = '[ 담당PM : 김더미 ]'
    ws['A6'] = '요청조직명'; ws['D6'] = '더미사업본부 경기사업단 평택운영팀'; ws['H6'] = '요청일자'; ws['K6'] = 46023
    ws['A7'] = '요청사유'; ws['D7'] = '2026년 최저임금 인상 및 단축상담 근무시간 조정'; ws['H7'] = '관리번호'; ws['K7'] = f'{CENTER}/2026년1월'
    ws['A9'] = 'Ⅱ. 직무Group'
    ws['A11'] = '직군명'; ws['D11'] = 'TM직군'
    ws['A12'] = '직무Group명'; ws['D12'] = CENTER; ws['H12'] = '직무Group정의'; ws['K12'] = '통신 고객상담'
    ws['A14'] = 'Ⅲ. 직무'
    for c, v in zip('ABEHS', ['구분', '직무명', '직무정의', '직무/직책수당', '정원']):
        ws[f'{c}15'] = v
    for i, (a, b) in enumerate([('조직장', '센터장'), (None, '실장'), (None, '팀장'), (None, '교육/QA'), ('조직원', '상담사')]):
        if a: ws[f'A{16 + i}'] = a
        ws[f'B{16 + i}'] = b; ws[f'E{16 + i}'] = f'{b} 업무'
    ws['A24'] = 'Ⅳ. 정원'
    ws['A25'] = '구분'; ws['B25'] = '직급'; ws['D25'] = '컨설턴트'; ws['B26'] = '호봉'; ws['A27'] = '정원'
    heads = [4, 30, 18, 12, 6, 10, 8, 4, 8, 4, 6, 5]
    for i, j in enumerate(jobs):
        ws.cell(26, 4 + i, f'{j[0]} {j[1]}'); ws.cell(27, 4 + i, heads[i])
    ws['A28'] = '구분'; ws['B28'] = '직급'; ws['D28'] = 'Staff'; ws['B29'] = '호봉'; ws['A30'] = '정원'
    for i, j in enumerate(leads):
        ws.cell(29, 4 + i, j[1]); ws.cell(30, 4 + i, [3, 4, 2, 1, 1][i])


def block_frame(ws, top, title, groups_row, names):
    ws.cell(top, 1, title); ws.cell(top, 18, '(단위 : 원)')
    ws.cell(top + 1, 1, '구분'); ws.cell(top + 1, 19, '비고')
    last = None
    for i, (g, n) in enumerate(names):
        if g != last: ws.cell(top + 1, 4 + i, g); last = g
        ws.cell(top + 2, 4 + i, n)
    for k, lab in enumerate(['기본급', '직책수당', '직무수당', '근속수당', '급식비', '통근비']):
        ws.cell(top + 3 + k, 1, lab)
    ws.cell(top + 9, 1, '통상임금소계'); ws.cell(top + 10, 1, '업적수당')
    for k, g in enumerate(GRADES):
        ws.cell(top + 10 + k, 3, g)
    ws.cell(top + 16, 1, '초과근무수당'); ws.cell(top + 17, 1, '야간근로수당'); ws.cell(top + 18, 1, '통신비'); ws.cell(top + 19, 1, '월 보수 합계')
    for k, lab in enumerate(['연차수당(1D)', '시간외수당(1H)', '휴일근로수당(1H)', '야근수당(1H)', '호칭별 연간 보수 합계']):
        ws.cell(top + 20 + k, 1, lab)
    ws.cell(top + 10, 19, '평가결과에 따름'); ws.cell(top + 19, 19, 'B등급 기준')


def money(ws, top, c, hours, mw, pos, duty, ten, meal, grades, ot, night):
    X = L(c)
    ws.cell(top + 3, c, f'={mw}*{hours}')
    if pos: ws.cell(top + 4, c, pos)
    if duty: ws.cell(top + 5, c, duty)
    if ten: ws.cell(top + 6, c, ten)
    ws.cell(top + 7, c, meal)
    ws.cell(top + 9, c, f'=SUM({X}{top + 3}:{X}{top + 8})')
    for k, v in enumerate(grades):
        ws.cell(top + 10 + k, c, v)
    ws.cell(top + 15, c, 0)
    ws.cell(top + 20, c, f'={X}{top + 9}/{hours}*8')
    ws.cell(top + 21, c, f'={X}{top + 9}/{hours}*1.5')
    ws.cell(top + 22, c, f'={X}{top + 9}/{hours}*1.5')
    ws.cell(top + 23, c, f'={X}{top + 9}/{hours}*0.5')
    if ot: ws.cell(top + 16, c, f'=ROUNDUP({X}{top + 21}*{ot},-1)')
    if night: ws.cell(top + 17, c, f'=ROUNDUP({X}{top + 23}*{night},-1)')
    ws.cell(top + 19, c, f'={X}{top + 9}+{X}{top + 12}+SUM({X}{top + 16}:{X}{top + 18})')
    ws.cell(top + 24, c, f'={X}{top + 19}*12')


def worktable(ws, when):
    ws['A94'] = 'Ⅵ. 예외 및 특례사항'
    ws['A95'] = '     1. 수습기간은 3개월 적용한다.'
    ws['A97'] = '     2. 무급사유 발생 시 통상임금 항목은 일할 계산한다.'
    ws['A99'] = '     3. 업적수당 : 성과평가 결과에 따라 직무별 차등 지급한다.'
    ws['A128'] = '  9.근무형태'
    for c, v in zip('ABCEFG', ['직무', ' 호봉', ' 근무일', ' 근무시간', ' 휴게시간', ' 비고']):
        ws[f'{c}129'] = v
    short = (('주 5일(월~금), 25시간(일 5H)', '09:00 ~ 14:30', '12:00 ~ 12:30') if when == 'before'
             else ('주 5일(월~금), 20시간(일 4H)', '09:00 ~ 13:00', '-'))
    rows = [('센터장, 실장', '센터장, 실장', '주 5일(월~금), 40시간(일 8H)', '09:00 ~ 18:00', '12:00 ~ 13:00'),
            ('팀장, 파트장, 강사', '팀장, 파트장, 강사', '주 5일(월~금), 40시간(일 8H)', '09:00 ~ 18:00', '12:00 ~ 13:00'),
            ('일반상담', '수습, 사원~책임', '주 5일(월~금), 40시간(일 8H)', '09:00 ~ 18:00', '12:00 ~ 13:00'),
            ('전문상담', '사원~선임', '주 5일(월~금), 40시간(일 8H)', '09:00 ~ 18:00', '12:00 ~ 13:00'),
            ('야간상담', '사원, 선임', '주 5일(일~목), 40시간(일 8H)', '22:00 ~ 07:00', '02:00 ~ 03:00'),
            ('단축상담', '사원', *short)]
    if when == 'after':
        rows.append(('채팅상담', '사원', '주 5일(월~금), 40시간(일 8H)', '10:00 ~ 19:00', '13:00 ~ 14:00'))
    for k, (a, b, d, t, r) in enumerate(rows):
        n = 130 + k
        ws[f'A{n}'] = a; ws[f'B{n}'] = b; ws[f'C{n}'] = d; ws[f'E{n}'] = t; ws[f'F{n}'] = r


def sheet(ws, when):
    mw = 10030 if when == 'before' else 10320
    meal = 150000 if when == 'before' else 200000
    jobs, leads = main_jobs(when), lead_jobs(when)
    header(ws, jobs, leads)
    block_frame(ws, 32, 'Ⅴ. 보수테이블', None, [(g, n) for g, n, *_ in jobs])
    for i, (g, n, hours, duty, ten, grades, ot, night) in enumerate(jobs):
        money(ws, 32, 4 + i, hours, mw, 0, duty, ten, meal, grades, ot, night)
    ws.cell(48, 19, '사원~책임 4H')
    block_frame(ws, 63, 'ㅇ 리딩그룹(스탭이상 ~ 센터장)', None, [(g, n) for g, n, *_ in leads])
    for i, (g, n, pos, duty, ten, ot) in enumerate(leads):
        money(ws, 63, 4 + i, 209, mw, pos, duty, ten, meal, [250000, 180000, 120000, 60000, 20000], ot, 0)
    ws.cell(79, 19, '파트장·팀장·강사·실장 5H')
    worktable(ws, when)


def main():
    wb = openpyxl.Workbook()
    wb.active.title = '급여테이블(변경전)'
    sheet(wb.active, 'before')
    sheet(wb.create_sheet('급여테이블(변경후)'), 'after')
    OUT.parent.mkdir(exist_ok=True)
    wb.save(OUT)
    recalc(OUT)
    print(OUT.name)


if __name__ == '__main__':
    main()
